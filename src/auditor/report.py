"""Write the audit results to a formatted Excel workbook."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "0C182B"
TEAL = "1F9E85"
LIGHT = "EEF3F7"
EUR = '#,##0.00 "€";-#,##0.00 "€"'
EUR0 = '#,##0 "€"'
PCT = "0.0%"
INT = "#,##0"

TITLE = Font(name="Calibri", size=18, bold=True, color=NAVY)
H2 = Font(name="Calibri", size=12, bold=True, color=NAVY)
LABEL = Font(name="Calibri", size=10, color="5B6B7F")
KPI = Font(name="Calibri", size=16, bold=True, color=NAVY)
THIN = Border(bottom=Side(style="thin", color="D5DEE8"))

EXCEPTION_LABELS = {
    "NO_SHIPMENT": "No matching shipment",
    "DUPLICATE": "Duplicate billing",
    "WEIGHT_OVERBILL": "Weight above chargeable",
    "RATE_ABOVE_CONTRACT": "Rate above contract",
    "WRONG_FUEL_PCT": "Fuel % above index",
    "UNAGREED_ACCESSORIAL": "Unagreed accessorial",
    "ARITHMETIC": "Total ≠ sum of charges",
}


def _write_table(ws, df: pd.DataFrame, top: int, left: int, name: str, formats: dict) -> int:
    """Write df as an Excel table; returns the next free row."""
    for j, col in enumerate(df.columns):
        h = ws.cell(row=top, column=left + j, value=col)
        # explicit header style so it also renders outside Excel (LibreOffice, previews)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=NAVY)
        h.alignment = Alignment(vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            c = ws.cell(row=top + i, column=left + j, value=v)
            if df.columns[j] in formats:
                c.number_format = formats[df.columns[j]]
    ref = f"{get_column_letter(left)}{top}:{get_column_letter(left + len(df.columns) - 1)}{top + len(df)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    return top + len(df) + 2


def build_summaries(r: pd.DataFrame) -> dict[str, pd.DataFrame]:
    ex = r[r.status == "EXCEPTION"]
    by_type = (
        ex.groupby("primary_exception")
        .agg(lines=("line_id", "count"), billed=("total_amount", "sum"),
             recoverable=("recoverable_eur", "sum"),
             no_impact=("recoverable_eur", lambda s: int((s == 0).sum())))
        .reindex([k for k in EXCEPTION_LABELS if k in set(ex.primary_exception)])
        .reset_index()
    )
    by_type.insert(1, "description", by_type.primary_exception.map(EXCEPTION_LABELS))
    by_type.columns = ["Exception", "Description", "Lines", "Billed €", "Recoverable €", "Lines with no € impact"]

    by_carrier = (
        r.groupby(["carrier_id", "carrier_name"])
        .agg(lines=("line_id", "count"),
             exceptions=("status", lambda s: int((s == "EXCEPTION").sum())),
             billed=("total_amount", "sum"),
             recoverable=("recoverable_eur", "sum"))
        .reset_index()
    )
    by_carrier["exception_rate"] = by_carrier.exceptions / by_carrier.lines
    by_carrier["recoverable_pct"] = by_carrier.recoverable / by_carrier.billed
    by_carrier = by_carrier.sort_values("recoverable", ascending=False)
    by_carrier.columns = ["Carrier", "Name", "Lines", "Exceptions", "Billed €",
                          "Recoverable €", "Exception rate", "Recoverable % of spend"]

    claims = (
        ex[ex.recoverable_eur > 0]
        .groupby(["carrier_id", "carrier_name", "invoice_id"])
        .agg(lines=("line_id", "count"), recoverable=("recoverable_eur", "sum"),
             reasons=("primary_exception", lambda s: ", ".join(sorted(set(s)))))
        .reset_index()
        .sort_values(["carrier_id", "invoice_id"])
    )
    claims.columns = ["Carrier", "Name", "Invoice", "Lines in dispute", "Claim €", "Reasons"]
    return dict(by_type=by_type, by_carrier=by_carrier, claims=claims)


def write_report(r: pd.DataFrame, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    s = build_summaries(r)
    wb = Workbook()

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Freight Invoice Audit"
    ws["B2"].font = TITLE
    period = f"{r.invoice_date.min():%d %b %Y} – {r.invoice_date.max():%d %b %Y}"
    ws["B3"] = f"Invoice lines dated {period} · synthetic data"
    ws["B3"].font = LABEL

    n = len(r)
    n_ex = int((r.status == "EXCEPTION").sum())
    kpis = [
        ("Invoice lines audited", n, INT),
        ("Carrier invoices", r.invoice_id.nunique(), INT),
        ("Total billed", r.total_amount.sum(), EUR0),
        ("Lines with exceptions", n_ex, INT),
        ("Clean-match rate", 1 - n_ex / n, PCT),
        ("Recoverable", r.recoverable_eur.sum(), EUR0),
        ("Recoverable % of spend", r.recoverable_eur.sum() / r.total_amount.sum(), PCT),
    ]
    for i, (label, val, fmt) in enumerate(kpis):
        col = 2 + (i % 4) * 2
        row = 5 + (i // 4) * 3
        ws.cell(row=row, column=col, value=label).font = LABEL
        c = ws.cell(row=row + 1, column=col, value=val)
        c.font = KPI
        c.number_format = fmt
        c.border = THIN

    ws["B12"] = "Exceptions by type"
    ws["B12"].font = H2
    nxt = _write_table(ws, s["by_type"], 13, 2, "ByType",
                       {"Billed €": EUR0, "Recoverable €": EUR0, "Lines": INT})

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Recoverable € by exception"
    chart.y_axis.numFmt = '#,##0 "€"'
    chart.legend = None
    data = Reference(ws, min_col=6, min_row=13, max_row=13 + len(s["by_type"]))
    cats = Reference(ws, min_col=3, min_row=14, max_row=13 + len(s["by_type"]))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.series[0].graphicalProperties.solidFill = TEAL
    chart.height, chart.width = 7.5, 16
    ws.add_chart(chart, "J12")

    ws.cell(row=nxt, column=2, value="By carrier").font = H2
    _write_table(ws, s["by_carrier"], nxt + 1, 2, "ByCarrier",
                 {"Billed €": EUR0, "Recoverable €": EUR0, "Exception rate": PCT,
                  "Recoverable % of spend": PCT, "Lines": INT})

    widths = {"A": 2, "B": 24, "C": 28, "D": 12, "E": 14, "F": 15, "G": 22, "H": 16, "I": 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # ---------------- Claims ----------------
    wc = wb.create_sheet("Carrier claims")
    wc.sheet_view.showGridLines = False
    wc["B2"] = "Amounts to dispute, per carrier invoice"
    wc["B2"].font = H2
    _write_table(wc, s["claims"], 4, 2, "Claims", {"Claim €": EUR})
    for col, w in {"A": 2, "B": 10, "C": 28, "D": 18, "E": 16, "F": 14, "G": 60}.items():
        wc.column_dimensions[col].width = w

    # ---------------- Exceptions detail ----------------
    we = wb.create_sheet("Exceptions")
    detail = r[r.status == "EXCEPTION"][[
        "line_id", "invoice_id", "carrier_id", "shipment_id", "invoice_date", "ship_date",
        "origin", "destination", "chargeable_weight_kg", "billed_weight_kg",
        "rate_per_kg", "contract_rate_per_kg", "fuel_pct", "contract_fuel_pct",
        "accessorial_codes", "unagreed_codes", "total_amount", "expected_total",
        "variance_eur", "recoverable_eur", "primary_exception", "exceptions",
    ]].copy()
    detail["invoice_date"] = detail.invoice_date.dt.date
    detail["ship_date"] = detail.ship_date.dt.date
    fmts = {c: EUR for c in ["total_amount", "expected_total", "variance_eur", "recoverable_eur"]}
    fmts.update({"fuel_pct": PCT, "contract_fuel_pct": PCT})
    _write_table(we, detail.fillna(""), 1, 1, "ExceptionLines", fmts)
    we.freeze_panes = "B2"
    for j, col in enumerate(detail.columns, start=1):
        we.column_dimensions[get_column_letter(j)].width = max(12, min(24, len(col) + 2))

    # ---------------- Method ----------------
    wm = wb.create_sheet("Method")
    wm.sheet_view.showGridLines = False
    lines = [
        ("How each line is checked", H2),
        ("1. Matched to its shipment record; no match = NO_SHIPMENT.", None),
        ("2. Same shipment billed more than once: later lines = DUPLICATE.", None),
        ("3. Chargeable weight = max(actual kg, m³ × 333). Billed weight > chargeable + 1% = WEIGHT_OVERBILL.", None),
        ("4. Rate per kg compared to the rate card for that lane and weight band = RATE_ABOVE_CONTRACT.", None),
        ("5. Fuel % compared to the carrier's index for the ship month (±0.05 pp) = WRONG_FUEL_PCT.", None),
        ("6. Each accessorial unit must be in the contract and on the shipment; amount must match = UNAGREED_ACCESSORIAL.", None),
        ("7. Total must equal base + fuel + accessorials (±0.02 €) = ARITHMETIC.", None),
        ("", None),
        ("Recoverable € = billed − re-rated contract price, only on exception lines that are overbilled.", None),
        ("Duplicates and lines with no shipment are fully recoverable.", None),
        ("", None),
        ("All carriers, lanes, rates and shipments are synthetic.", LABEL),
    ]
    for i, (text, font) in enumerate(lines, start=2):
        c = wm.cell(row=i, column=2, value=text)
        if font:
            c.font = font
        c.alignment = Alignment(wrap_text=False)
    wm.column_dimensions["B"].width = 110

    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(path)
    return path
